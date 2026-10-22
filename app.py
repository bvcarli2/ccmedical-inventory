import os,sqlite3,re,sys
from datetime import datetime, timezone, date
from pathlib import Path
from flask import Flask,request,render_template_string,redirect,url_for,flash,session
from openpyxl import load_workbook
BASE=Path(__file__).resolve().parent
DB=BASE/'inventory.db'
app=Flask(__name__)
# Never expose the local prototype to the internet without explicit secure mode.
app.secret_key=os.environ.get('CCM_SECRET','local-prototype-only')
app.config.update(
 SESSION_COOKIE_HTTPONLY=True,
 SESSION_COOKIE_SAMESITE='Lax',
 SESSION_COOKIE_SECURE=os.environ.get('CCM_ENABLE_EMPLOYEE_AUTH')=='1',
)
if os.environ.get('RENDER'):
 raise RuntimeError('Deployment blocked until PostgreSQL migration, role authorization and CSRF protection are complete.')
if os.environ.get('CCM_ENABLE_EMPLOYEE_AUTH')=='1' and (
 not os.environ.get('CCM_SECRET') or os.environ.get('CCM_SECRET')=='local-prototype-only'
):
 raise RuntimeError('A unique CCM_SECRET is required for employee authentication.')

SCHEMA='''CREATE TABLE IF NOT EXISTS units(id INTEGER PRIMARY KEY,source_sheet TEXT,source_row INTEGER,manufacturer TEXT,product_raw TEXT,product_key TEXT,serial TEXT,cost REAL,dom TEXT,notes TEXT,location_raw TEXT,physical_location TEXT,status TEXT,condition TEXT,review TEXT, UNIQUE(source_sheet,source_row)); CREATE TABLE IF NOT EXISTS movements(id INTEGER PRIMARY KEY,unit_id INTEGER,old_location TEXT,new_location TEXT,created_at TEXT DEFAULT CURRENT_TIMESTAMP,note TEXT); CREATE TABLE IF NOT EXISTS repair_orders(id INTEGER PRIMARY KEY,unit_id INTEGER,repair_type TEXT,status TEXT DEFAULT 'Open',notes TEXT,created_at TEXT DEFAULT CURRENT_TIMESTAMP); CREATE TABLE IF NOT EXISTS integration_links(id INTEGER PRIMARY KEY,unit_id INTEGER,system TEXT,external_id TEXT,sync_state TEXT DEFAULT 'not_synced');'''
def conn():
 c=sqlite3.connect(DB);c.row_factory=sqlite3.Row;return c
def norm(s):return re.sub(r'\s+',' ',str(s or '').strip()).upper()
def classify(loc):
 s=norm(loc)
 if s in ('SOLD','S0LD','CLOSED'):return 'Sold archive',None
 if s=='RENTAL':return 'Rental',None
 if s in ('LOAN','LOANER'):return 'On loan',None
 if s in ('PENDING','N'):return 'Pending review',None
 if re.match(r'^[A-Z]\d+(\s*/\s*[A-Z]\d+)?$',s):return 'Active',s.replace(' ','')
 return 'Needs review',None

def import_file(path):
 w=load_workbook(path,read_only=True,data_only=True);s=w['Location'];c=conn();c.executescript(SCHEMA);added=0
 for rownum,row in enumerate(s.iter_rows(min_row=3,values_only=True),start=3):
  loc,prod,serial,cost,dom,notes=(list(row)+[None]*6)[:6]
  if not prod or not str(prod).strip():continue
  status,physical=classify(loc);product=norm(prod);serial=str(serial).strip() if serial is not None else ''
  issues=[]
  if not serial:issues.append('Missing serial')
  if status=='Needs review':issues.append('Unrecognized location/status')
  if status=='Pending review':issues.append('Pending status')
  if cost is None or str(cost).strip()=='':issues.append('Missing cost')
  try:amount=float(cost) if cost is not None and str(cost).strip() else None
  except (TypeError,ValueError):amount=None;issues.append('Invalid cost')
  if amount==0:issues.append('Zero cost (verify)')
  if 'NEW' in product:condition='New (inferred)'
  elif 'REFURB' in product:condition='Refurbished (inferred)'
  else:condition='Unspecified'
  c.execute('INSERT OR IGNORE INTO units(source_sheet,source_row,manufacturer,product_raw,product_key,serial,cost,dom,notes,location_raw,physical_location,status,condition,review) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)',('Location',rownum,None,str(prod).strip(),product,serial,amount,str(dom or ''),str(notes or ''),str(loc or ''),physical,status,condition,'; '.join(issues)))
  added+=c.execute('SELECT changes()').fetchone()[0]
 # Flag duplicates for review, not merge
 duplicates=c.execute("SELECT UPPER(TRIM(serial)) serial,COUNT(*) n FROM units WHERE TRIM(serial)<>'' GROUP BY UPPER(TRIM(serial)) HAVING COUNT(*)>1").fetchall()
 for d in duplicates:
  c.execute("UPDATE units SET review=TRIM(COALESCE(review,'')||'; Duplicate serial (verify)', '; ') WHERE UPPER(TRIM(serial))=?",(d['serial'],))
 c.commit();c.close();return added,len(duplicates)
REPAIR_STAGES=['Open Repairs','Evaled','In Repair','QC','Ready to Ship','Completed']
REPAIR_SCHEMA="""CREATE TABLE IF NOT EXISTS customer_repairs (
 id INTEGER PRIMARY KEY, rma TEXT NOT NULL UNIQUE COLLATE NOCASE, customer TEXT NOT NULL,
 manufacturer TEXT, model TEXT, part_number TEXT, serial TEXT, location TEXT,
 received_date TEXT NOT NULL, technician TEXT, notes TEXT,
 stage TEXT NOT NULL DEFAULT 'Open Repairs', stage_started_at TEXT NOT NULL,
 created_at TEXT NOT NULL, completed_at TEXT, updated_at TEXT NOT NULL);
 CREATE TABLE IF NOT EXISTS customer_repair_history (
 id INTEGER PRIMARY KEY, repair_id INTEGER NOT NULL REFERENCES customer_repairs(id),
 old_stage TEXT, new_stage TEXT NOT NULL, changed_at TEXT NOT NULL, note TEXT);
 CREATE INDEX IF NOT EXISTS idx_repair_history_repair ON customer_repair_history(repair_id);
"""
def now_iso():return datetime.now(timezone.utc).isoformat(timespec='seconds')
def days_since(ts):
 try:return max(0,(datetime.now(timezone.utc).date()-datetime.fromisoformat(ts.replace('Z','+00:00')).date()).days)
 except (ValueError,AttributeError):return 0
def age_class(days,stage):
 if stage=='Completed':return 'complete'
 return 'red' if days>7 else 'yellow' if days>3 else 'normal'
def repair_context(r):
 d=dict(r)
 d['stage_days']=days_since(d['stage_started_at']) if d['stage']!='Completed' else 0
 d['total_days']=max(0,((datetime.fromisoformat(d['completed_at']).date() if d['completed_at'] else datetime.now(timezone.utc).date())-date.fromisoformat(d['received_date'])).days)
 d['alert']=age_class(d['stage_days'],d['stage'])
 return d
# New inventory entry: a single shared, never-reset sequence across CC, RP and CR.
ENTRY_SCHEMA="""CREATE TABLE IF NOT EXISTS inventory_id_sequence (
 seq INTEGER PRIMARY KEY AUTOINCREMENT, prefix TEXT NOT NULL,
 generated_id TEXT UNIQUE, created_at TEXT NOT NULL);
 CREATE TABLE IF NOT EXISTS core_items (
 id INTEGER PRIMARY KEY, inventory_code TEXT NOT NULL UNIQUE,
 manufacturer TEXT, product_raw TEXT NOT NULL, serial TEXT, location TEXT,
 supplier TEXT, original_cost REAL NOT NULL DEFAULT 0,
 stage TEXT NOT NULL DEFAULT 'Needs Repair', notes TEXT,
 qc_result TEXT, qc_technician TEXT, qc_date TEXT, qc_notes TEXT,
 converted_unit_id INTEGER, created_at TEXT NOT NULL);
 CREATE TABLE IF NOT EXISTS core_transactions (
 id INTEGER PRIMARY KEY, core_id INTEGER NOT NULL,
 performed_at TEXT NOT NULL, technician TEXT, description TEXT NOT NULL,
 labor_method TEXT NOT NULL DEFAULT 'manual', labor_hours REAL NOT NULL DEFAULT 0,
 labor_rate REAL NOT NULL DEFAULT 0, labor_cost REAL NOT NULL DEFAULT 0,
 parts_cost REAL NOT NULL DEFAULT 0, other_cost REAL NOT NULL DEFAULT 0);
 CREATE TABLE IF NOT EXISTS receiving_batches (
 id INTEGER PRIMARY KEY, created_at TEXT NOT NULL, item_count INTEGER NOT NULL);
"""

def ensure_entry_schema(c):
 c.executescript(ENTRY_SCHEMA)
 cols={x['name'] for x in c.execute('PRAGMA table_info(units)')}
 for name,typ in [('inventory_code','TEXT'),('supplier','TEXT'),('received_date','TEXT'),('entry_batch_id','INTEGER')]:
  if name not in cols:c.execute('ALTER TABLE units ADD COLUMN '+name+' '+typ)
 cols={x['name'] for x in c.execute('PRAGMA table_info(customer_repairs)')}
 if 'inventory_code' not in cols:c.execute('ALTER TABLE customer_repairs ADD COLUMN inventory_code TEXT')
 # Existing RMA records are retained; IDs are assigned only to new entries.
 c.execute('CREATE UNIQUE INDEX IF NOT EXISTS idx_unit_inventory_code ON units(inventory_code)')
 c.execute('CREATE UNIQUE INDEX IF NOT EXISTS idx_repair_inventory_code ON customer_repairs(inventory_code)')

def new_inventory_code(c,prefix,received_date=None):
 if prefix not in ('CC','RP','CR'):raise ValueError('Unknown inventory prefix')
 today=date.today().strftime('%m%d%y') # Entry date, not historical received date
 now=now_iso()
 cur=c.execute('INSERT INTO inventory_id_sequence(prefix,created_at) VALUES (?,?)',(prefix,now))
 seq=cur.lastrowid
 code=f'{prefix}-{today}-{seq:03d}'
 c.execute('UPDATE inventory_id_sequence SET generated_id=? WHERE seq=?',(code,seq))
 return code

def amount(value,label,required=False):
 raw=str(value or '').strip()
 if not raw:return 0.0 if not required else None
 try:
  n=float(raw)
  if not (0<=n<1e12):raise ValueError()
  return round(n,2)
 except ValueError:raise ValueError(label+' must be a nonnegative number')

def entry_row(f):
 return {k:str(f.get(k,'') or '').strip() for k in ('type','manufacturer','product','serial','location','supplier','cost','condition','notes','customer','rma','technician','received_date')}

@app.route('/inventory/new',methods=['GET','POST'])
def inventory_new():
 if request.method=='GET':return render_template_string(ENTRY_TEMPLATE,today=date.today().isoformat())
 # Multiple-row entry is atomic: invalid rows do not result in partial receiving.
 f=request.form;fields=('type','manufacturer','product','serial','location','supplier','cost','condition','notes','customer','rma','technician','received_date')
 length=max((len(f.getlist(k+'[]')) for k in fields),default=0)
 rows=[]
 for i in range(length):
  row={k:(f.getlist(k+'[]')[i].strip() if i<len(f.getlist(k+'[]')) else '') for k in fields}
  if not any(row.values()):continue
  if not row['type'] or row['type'] not in ('CC','CR','RP'):return 'Select CC, CR or RP for every row',400
  if not row['product']:return 'Model / product is required for every row',400
  if row['type']=='RP' and (not row['customer'] or not row['rma']):return 'Customer and unique RMA are required for RP items',400
  try:
   if row['received_date']:date.fromisoformat(row['received_date'])
   row['cost_value']=amount(row['cost'],'Purchase cost')
  except ValueError as e:return str(e),400
  rows.append(row)
 if not rows:return 'Enter at least one item',400
 c=conn()
 try:
  ensure_entry_schema(c)
  c.execute('BEGIN IMMEDIATE')
  batch=c.execute('INSERT INTO receiving_batches(created_at,item_count) VALUES (?,?)',(now_iso(),len(rows))).lastrowid
  codes=[]
  for row in rows:
   typ=row['type'];code=new_inventory_code(c,typ);codes.append(code)
   received=row['received_date'] or date.today().isoformat()
   if typ=='CC':
    c.execute("""INSERT INTO units(manufacturer,product_raw,product_key,serial,cost,dom,notes,location_raw,physical_location,status,condition,review,inventory_code,supplier,received_date,entry_batch_id)
     VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
     (row['manufacturer'],row['product'],norm(row['product']),row['serial'],row['cost_value'],'',row['notes'],row['location'],row['location'],'Active',row['condition'] or 'Unspecified','',code,row['supplier'],received,batch))
   elif typ=='CR':
    c.execute("""INSERT INTO core_items(inventory_code,manufacturer,product_raw,serial,location,supplier,original_cost,notes,created_at)
     VALUES (?,?,?,?,?,?,?,?,?)""",(code,row['manufacturer'],row['product'],row['serial'],row['location'],row['supplier'],row['cost_value'],row['notes'],now_iso()))
   else:
    now=now_iso()
    cur=c.execute("""INSERT INTO customer_repairs(rma,customer,manufacturer,model,part_number,serial,location,received_date,technician,notes,stage,stage_started_at,created_at,updated_at,inventory_code)
     VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(row['rma'],row['customer'],row['manufacturer'],row['product'],'',row['serial'],row['location'],received,row['technician'],row['notes'],'Open Repairs',now,now,now,code))
    c.execute('INSERT INTO customer_repair_history(repair_id,old_stage,new_stage,changed_at,note) VALUES (?,?,?,?,?)',(cur.lastrowid,None,'Open Repairs',now,'Received through New Inventory Entry'))
  c.commit()
 except (sqlite3.Error,ValueError) as e:
  c.rollback();c.close();return 'Receiving not saved: '+str(e),400
 c.close();flash('Saved '+str(len(codes))+' items: '+', '.join(codes));return redirect(url_for('inventory_new'))

@app.route('/cores')
def cores():
 c=conn();ensure_entry_schema(c)
 rows=c.execute("""SELECT c.*,COALESCE((SELECT SUM(labor_cost+parts_cost+other_cost) FROM core_transactions WHERE core_id=c.id),0) repair_cost FROM core_items c ORDER BY c.id DESC""").fetchall();c.close()
 return render_template_string(CORE_TEMPLATE,rows=rows,detail=None,transactions=[],total=0)

@app.route('/cores/<int:cid>',methods=['GET','POST'])
def core_detail(cid):
 c=conn();ensure_entry_schema(c);core=c.execute('SELECT * FROM core_items WHERE id=?',(cid,)).fetchone()
 if core is None:c.close();return 'Core not found',404
 if request.method=='POST':
  f=request.form;action=f.get('action','')
  try:
   c.execute('BEGIN IMMEDIATE')
   if core['stage']=='Converted to Stock':raise ValueError('Converted cores cannot be modified')
   if action=='transaction':
    method=f.get('labor_method','manual');hours=amount(f.get('hours'),'Hours');rate=amount(f.get('rate'),'Hourly rate')
    labor=amount(f.get('labor'),'Labor cost') if method=='manual' else round(hours*rate,2)
    parts=amount(f.get('parts'),'Parts cost');other=amount(f.get('other'),'Other cost')
    if not f.get('description','').strip():raise ValueError('Transaction description required')
    c.execute("""INSERT INTO core_transactions(core_id,performed_at,technician,description,labor_method,labor_hours,labor_rate,labor_cost,parts_cost,other_cost) VALUES (?,?,?,?,?,?,?,?,?,?)""",(cid,now_iso(),f.get('technician','').strip(),f['description'].strip(),method,hours,rate,labor,parts,other))
   elif action=='stage':
    stage=f.get('stage')
    if stage not in ('Needs Repair','In Repair','QC'):raise ValueError('Invalid stage change')
    c.execute('UPDATE core_items SET stage=?,qc_result=NULL WHERE id=?',(stage,cid))
   elif action=='qc':
    result=f.get('result');tech=f.get('technician','').strip()
    if core['stage']!='QC':raise ValueError('Move the core to QC first')
    if result not in ('Passed','Failed') or not tech:raise ValueError('QC result and technician required')
    c.execute('UPDATE core_items SET qc_result=?,qc_technician=?,qc_date=?,qc_notes=?,stage=? WHERE id=?',(result,tech,now_iso(),f.get('notes','').strip(),'Ready for Stock' if result=='Passed' else 'In Repair',cid))
   elif action=='convert':
    if core['stage']!='Ready for Stock' or core['qc_result']!='Passed':raise ValueError('Passing QC is required before conversion')
    repair_cost=c.execute('SELECT COALESCE(SUM(labor_cost+parts_cost+other_cost),0) FROM core_transactions WHERE core_id=?',(cid,)).fetchone()[0]
    code=new_inventory_code(c,'CC');total=round(core['original_cost']+repair_cost,2)
    cur=c.execute("""INSERT INTO units(manufacturer,product_raw,product_key,serial,cost,notes,location_raw,physical_location,status,condition,review,inventory_code,supplier,received_date)
     VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(core['manufacturer'],core['product_raw'],norm(core['product_raw']),core['serial'],total,'Converted from '+core['inventory_code']+'; '+(core['notes'] or ''),core['location'],core['location'],'Active','Refurbished','',code,core['supplier'],date.today().isoformat()))
    c.execute('UPDATE core_items SET stage=?,converted_unit_id=? WHERE id=?',('Converted to Stock',cur.lastrowid,cid))
    flash('Converted to '+code+' with total cost $'+format(total,',.2f'))
   else:raise ValueError('Unknown action')
   c.commit()
  except (ValueError,sqlite3.Error) as e:c.rollback();c.close();return str(e),400
  c.close();return redirect(url_for('core_detail',cid=cid))
 core=c.execute('SELECT * FROM core_items WHERE id=?',(cid,)).fetchone()
 tx=c.execute('SELECT * FROM core_transactions WHERE core_id=? ORDER BY id DESC',(cid,)).fetchall()
 total=core['original_cost']+sum(t['labor_cost']+t['parts_cost']+t['other_cost'] for t in tx)
 c.close();return render_template_string(CORE_TEMPLATE,detail=core,transactions=tx,total=total,rows=[])

TEMPLATE=(BASE/'dashboard.html').read_text(encoding='utf-8')
REPAIR_TEMPLATE=(BASE/'repairs.html').read_text(encoding='utf-8')
ENTRY_TEMPLATE=(BASE/'inventory_entry.html').read_text(encoding='utf-8')
CORE_TEMPLATE=(BASE/'cores.html').read_text(encoding='utf-8')
@app.route('/repairs')
def repairs():
 c=conn();c.executescript(REPAIR_SCHEMA)
 stage=request.args.get('stage','');q=request.args.get('q','').strip()
 if stage and stage not in REPAIR_STAGES:stage=''
 filters=[];args=[]
 if stage:filters.append('stage=?');args.append(stage)
 if q:
  filters.append('(rma LIKE ? OR customer LIKE ? OR serial LIKE ? OR location LIKE ? OR model LIKE ? OR part_number LIKE ?)');args.extend(['%'+q+'%']*6)
 where=(' WHERE '+' AND '.join(filters)) if filters else ''
 rows=c.execute('SELECT * FROM customer_repairs'+where+' ORDER BY CASE WHEN stage="Completed" THEN 1 ELSE 0 END, stage_started_at ASC, id DESC',args).fetchall()
 counts={r['stage']:r['n'] for r in c.execute('SELECT stage,COUNT(*) n FROM customer_repairs GROUP BY stage')}
 c.close()
 return render_template_string(REPAIR_TEMPLATE,mode='list',rows=[repair_context(r) for r in rows],counts=counts,stages=REPAIR_STAGES,selected=stage,q=q)
@app.route('/repairs/new',methods=['GET','POST'])
def repair_new():
 if request.method=='POST':
  f=request.form;rma=f.get('rma','').strip();customer=f.get('customer','').strip();received=f.get('received_date','').strip()
  if not rma or not customer:flash('RMA number and customer are required.');return redirect(url_for('repair_new'))
  try:date.fromisoformat(received)
  except ValueError:flash('Enter a valid received date.');return redirect(url_for('repair_new'))
  now=now_iso();c=conn();c.executescript(REPAIR_SCHEMA);ensure_entry_schema(c)
  try:
   c.execute('BEGIN IMMEDIATE')
   rp_code=new_inventory_code(c,'RP')
   cur=c.execute('INSERT INTO customer_repairs(rma,customer,manufacturer,model,part_number,serial,location,received_date,technician,notes,stage,stage_started_at,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
    (rma,customer,f.get('manufacturer','').strip(),f.get('model','').strip(),f.get('part_number','').strip(),f.get('serial','').strip(),f.get('location','').strip(),received,f.get('technician','').strip(),f.get('notes','').strip(),'Open Repairs',now,now,now))
   c.execute('UPDATE customer_repairs SET inventory_code=? WHERE id=?',(rp_code,cur.lastrowid))
   rid=cur.lastrowid
   c.execute('INSERT INTO customer_repair_history(repair_id,old_stage,new_stage,changed_at,note) VALUES (?,?,?,?,?)',(rid,None,'Open Repairs',now,'RMA created'))
   c.commit()
  except sqlite3.IntegrityError:
   c.rollback();c.close();flash('This RMA number already exists. Each item needs its own unique RMA.');return redirect(url_for('repair_new'))
  c.close();flash('RMA created in Open Repairs');return redirect(url_for('repair_detail',rid=rid))
 return render_template_string(REPAIR_TEMPLATE,mode='new',stages=REPAIR_STAGES,today=date.today().isoformat())
@app.route('/repairs/<int:rid>')
def repair_detail(rid):
 c=conn();c.executescript(REPAIR_SCHEMA);r=c.execute('SELECT * FROM customer_repairs WHERE id=?',(rid,)).fetchone()
 if r is None:c.close();return 'RMA not found',404
 history=c.execute('SELECT * FROM customer_repair_history WHERE repair_id=? ORDER BY id DESC',(rid,)).fetchall();c.close()
 return render_template_string(REPAIR_TEMPLATE,mode='detail',r=repair_context(r),history=history,stages=REPAIR_STAGES)
@app.post('/repairs/<int:rid>/update')
def repair_update(rid):
 f=request.form;new_stage=f.get('stage','')
 if new_stage not in REPAIR_STAGES:return 'Invalid repair stage',400
 c=conn();c.executescript(REPAIR_SCHEMA);r=c.execute('SELECT * FROM customer_repairs WHERE id=?',(rid,)).fetchone()
 if r is None:c.close();return 'RMA not found',404
 now=now_iso();stage_changed=new_stage!=r['stage']
 # A completed RMA can be reopened; completion timestamp is cleared in that case.
 complete=now if stage_changed and new_stage=='Completed' else (None if stage_changed else r['completed_at'])
 c.execute('UPDATE customer_repairs SET customer=?,manufacturer=?,model=?,part_number=?,serial=?,location=?,technician=?,notes=?,stage=?,stage_started_at=?,completed_at=?,updated_at=? WHERE id=?',
 (f.get('customer','').strip() or r['customer'],f.get('manufacturer','').strip(),f.get('model','').strip(),f.get('part_number','').strip(),f.get('serial','').strip(),f.get('location','').strip(),f.get('technician','').strip(),f.get('notes','').strip(),new_stage,now if stage_changed else r['stage_started_at'],complete,now,rid))
 if stage_changed:c.execute('INSERT INTO customer_repair_history(repair_id,old_stage,new_stage,changed_at,note) VALUES (?,?,?,?,?)',(rid,r['stage'],new_stage,now,f.get('change_note','').strip()))
 c.commit();c.close();flash('RMA updated');return redirect(url_for('repair_detail',rid=rid))

@app.route('/')
def index():
 view=request.args.get('view','Active');q=request.args.get('q','').strip();page=max(1,int(request.args.get('page','1')));c=conn();where=[];params=[]
 if view=='Active':where.append("status IN ('Active','Rental')")
 elif view!='All':where.append('status=?');params.append(view)
 if q:where.append('(product_raw LIKE ? OR serial LIKE ? OR notes LIKE ? OR location_raw LIKE ?)');params.extend(['%'+q+'%']*4)
 clause=' WHERE '+' AND '.join(where) if where else ''
 stats=c.execute('SELECT COUNT(*) n,COUNT(cost) priced,COALESCE(SUM(cost),0) val FROM units'+clause,params).fetchone();overview={r['status']:r['n'] for r in c.execute('SELECT status, COUNT(*) n FROM units GROUP BY status')};all_count=sum(overview.values());rows=c.execute('SELECT * FROM units'+clause+' ORDER BY id LIMIT 300 OFFSET ?',params+[(page-1)*300]).fetchall();c.close()
 return render_template_string(TEMPLATE,view=view,q=q,page=page,total=stats['n'],priced=stats['priced'],value=stats['val'],units=rows,unit=None,overview=overview,all_count=all_count)
@app.route('/unit/<int:uid>')
def unit(uid):
 c=conn();u=c.execute('SELECT * FROM units WHERE id=?',(uid,)).fetchone();c.close()
 if not u:return 'Not found',404
 return render_template_string(TEMPLATE,view=u['status'],unit=u,overview={},all_count=0)
@app.post('/transfer/<int:uid>')
def transfer(uid):
 new=norm(request.form.get('location'));status,physical=classify(new)
 if status!='Active':flash('Transfer requires a recognized physical location');return redirect(url_for('unit',uid=uid))
 c=conn();u=c.execute('SELECT physical_location,status FROM units WHERE id=?',(uid,)).fetchone()
 if not u or u['status']=='Sold archive':c.close();return 'Sold records cannot be transferred',400
 c.execute('INSERT INTO movements(unit_id,old_location,new_location,note) VALUES (?,?,?,?)',(uid,u['physical_location'],physical,request.form.get('note','')));c.execute('UPDATE units SET physical_location=?,status=? WHERE id=?',(physical,'Active',uid));c.commit();c.close();flash('Transfer recorded');return redirect(url_for('unit',uid=uid))
@app.post('/repair/<int:uid>')
def repair(uid):
 c=conn();u=c.execute('SELECT status FROM units WHERE id=?',(uid,)).fetchone()
 if not u or u['status']=='Sold archive':c.close();return 'Sold records cannot receive new repairs in this prototype',400
 c.execute('INSERT INTO repair_orders(unit_id,repair_type,notes) VALUES (?,?,?)',(uid,request.form.get('type'),request.form.get('notes')));c.commit();c.close();flash('Repair record created');return redirect(url_for('unit',uid=uid))
# Opt-in local employee security test. The Render deployment remains blocked
# until PostgreSQL migration and write-route security reviews are complete.
if os.environ.get('CCM_ENABLE_EMPLOYEE_AUTH')=='1':
 from employee_auth import register_employee_auth
 register_employee_auth(app,conn)

# Local-only PostgreSQL staging dispatch. Existing SQLite paths stay unchanged.
# Render startup is intentionally blocked above until security is complete.
if os.environ.get('CCM_PG_STAGING')=='1':
 if os.environ.get('CCM_ENABLE_EMPLOYEE_AUTH')=='1':
  raise RuntimeError('PostgreSQL staging bridge is incompatible with SQLite employee auth')
 if not os.environ.get('CCM_STAGING_DATABASE_URL'):
  raise RuntimeError('CCM_STAGING_DATABASE_URL is required for PostgreSQL staging')
 from postgres_staging_routes import install_staging_routes
 install_staging_routes(app,norm,classify,amount)

from product_database import register_product_routes
# PostgreSQL read-only dashboard and reporting routes for local staging.
if os.environ.get('CCM_PG_STAGING')=='1':
 from postgres_dashboard import install_dashboard_routes
 install_dashboard_routes(app,os.environ['CCM_STAGING_DATABASE_URL'],
                          TEMPLATE,CORE_TEMPLATE,REPAIR_TEMPLATE,
                          REPAIR_STAGES,repair_context)

if os.environ.get('CCM_PG_STAGING')=='1':
 from postgres_backend import PGConnection
 def product_pg_conn():
  return PGConnection(os.environ['CCM_STAGING_DATABASE_URL'],strict_inserts=False)
 register_product_routes(app,product_pg_conn,BASE,postgres=True)
else:
 register_product_routes(app,conn,BASE)

if __name__=='__main__':
 if len(sys.argv)>1 and sys.argv[1]=='import':
  print('Imported records, duplicate serial groups:',import_file(sys.argv[2]));sys.exit()
 c=conn();c.executescript(SCHEMA);c.executescript(REPAIR_SCHEMA);ensure_entry_schema(c);c.execute("UPDATE units SET status='Rental' WHERE UPPER(TRIM(location_raw))='RENTAL' AND status='On loan'");c.commit();c.close();app.run(host='127.0.0.1',port=5000,debug=False)
